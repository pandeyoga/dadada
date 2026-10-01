"""Backend tests for Leads master import + full-session raw import.

Coverage:
- Template.xlsx contains a 'Leads' sheet with expected columns.
- POST /api/data-mgmt/import (upsert) with template dry-run/commit → leads inserted,
  stage=acquisition, phone normalized (+62), assigned_to set, score numeric.
- Re-import same file → update/skip, no duplicates.
- Row without phone → error.
- POST /api/data-mgmt/full/sessions with raw file (/tmp/leads.xlsx) → 77 insert plan.
- POST /sessions/{sid}/commit → totals.inserted=77, leads persisted with defaults.
- Cleanup test leads.
"""
import io
import os
import pytest
import requests
import openpyxl

BASE_URL = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")
SUPERADMIN = ("superadmin@sipro.co.id", "Sipro#2026")
RAW_FILE = "/tmp/leads.xlsx"

# 3 unique test phones — must be easy to cleanup afterwards.
TEST_PHONES_UNIQUE = ["081299999901", "081299999902", "081299999903"]
TEST_NAMES = ["TEST_Andi", "TEST_Budi", "TEST_Citra"]


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{BASE_URL}/api/auth/login",
                      json={"email": SUPERADMIN[0], "password": SUPERADMIN[1]}, timeout=30)
    assert r.status_code == 200, f"login failed: {r.status_code} {r.text}"
    tok = r.json().get("token") or r.json().get("access_token")
    assert tok
    return tok


@pytest.fixture(scope="module")
def headers(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def template_bytes(headers):
    r = requests.get(f"{BASE_URL}/api/data-mgmt/template.xlsx?with_example=true",
                     headers=headers, timeout=60)
    assert r.status_code == 200, r.text
    return r.content


def _load(content):
    return openpyxl.load_workbook(io.BytesIO(content))


# --- Template shape --------------------------------------------------------
def test_template_has_leads_sheet(template_bytes):
    wb = _load(template_bytes)
    assert "Leads" in wb.sheetnames, f"sheets={wb.sheetnames}"
    ws = wb["Leads"]
    keys = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    expected = {"name", "phone", "email", "source", "campaign",
                "interest_unit_type", "assigned_to", "notes"}
    missing = expected - set(keys)
    assert not missing, f"missing columns: {missing}; got {keys}"


def _build_leads_only_upload(template_bytes, rows):
    """Return an .xlsx bytes with only the Leads sheet populated (other sheets
    stripped of data rows). rows = list of dicts."""
    wb = _load(template_bytes)
    # Wipe data rows (keep row1=keys, row2=labels) on every entity sheet
    for name in list(wb.sheetnames):
        ws = wb[name]
        if ws.max_row > 2:
            ws.delete_rows(3, ws.max_row - 2)
    ws = wb["Leads"]
    keys = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    for row in rows:
        ws.append([row.get(k) for k in keys])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- Template import (dry + commit) ----------------------------------------
def test_import_template_leads_dry_and_commit(headers, template_bytes):
    payload = _build_leads_only_upload(template_bytes, [
        {"name": TEST_NAMES[0], "phone": TEST_PHONES_UNIQUE[0]},
        {"name": TEST_NAMES[1], "phone": TEST_PHONES_UNIQUE[1]},
        {"name": TEST_NAMES[2], "phone": TEST_PHONES_UNIQUE[2],
         "assigned_to": "sales@sipro.co.id"},
    ])
    files = {"file": ("leads.xlsx", payload,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    # DRY RUN
    r = requests.post(f"{BASE_URL}/api/data-mgmt/import",
                      files=files, data={"mode": "upsert", "dry_run": "true"},
                      headers=headers, timeout=60)
    assert r.status_code == 200, r.text
    js = r.json()
    leads_ent = next((e for e in js["entities"] if e["key"] == "leads"), None)
    assert leads_ent is not None, f"no leads entity in report: {js}"
    assert leads_ent["insert"] == 3, leads_ent
    assert leads_ent["error"] == 0, leads_ent

    # COMMIT
    files = {"file": ("leads.xlsx", payload,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = requests.post(f"{BASE_URL}/api/data-mgmt/import",
                      files=files, data={"mode": "upsert", "dry_run": "false"},
                      headers=headers, timeout=60)
    assert r.status_code == 200, r.text
    js = r.json()
    leads_ent = next(e for e in js["entities"] if e["key"] == "leads")
    assert leads_ent["insert"] == 3, leads_ent

    # Verify persistence via /api/leads (search by TEST prefix)
    r = requests.get(f"{BASE_URL}/api/leads?q=TEST_&limit=50", headers=headers, timeout=30)
    assert r.status_code == 200, r.text
    body = r.json()
    items = body.get("data") or body.get("items") or (body if isinstance(body, list) else [])
    by_name = {i.get("name"): i for i in items}
    for name, phone in zip(TEST_NAMES, TEST_PHONES_UNIQUE):
        assert name in by_name, f"missing {name}"
        lead = by_name[name]
        assert lead.get("stage") == "acquisition", lead
        assert (lead.get("phone") or "").startswith("+62"), lead
        assert lead.get("assigned_to"), f"no assignee: {lead}"
        assert isinstance(lead.get("score"), (int, float)), lead


def test_reimport_same_leads_no_duplicate(headers, template_bytes):
    payload = _build_leads_only_upload(template_bytes, [
        {"name": TEST_NAMES[0], "phone": TEST_PHONES_UNIQUE[0]},
        {"name": TEST_NAMES[1], "phone": TEST_PHONES_UNIQUE[1]},
        {"name": TEST_NAMES[2], "phone": TEST_PHONES_UNIQUE[2],
         "assigned_to": "sales@sipro.co.id"},
    ])
    files = {"file": ("leads.xlsx", payload,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = requests.post(f"{BASE_URL}/api/data-mgmt/import",
                      files=files, data={"mode": "upsert", "dry_run": "false"},
                      headers=headers, timeout=60)
    assert r.status_code == 200
    leads_ent = next(e for e in r.json()["entities"] if e["key"] == "leads")
    # No inserts on second run; either update or skip
    assert leads_ent["insert"] == 0, leads_ent
    assert leads_ent["update"] + leads_ent["skip"] == 3, leads_ent


def test_leads_without_phone_reports_error(headers, template_bytes):
    payload = _build_leads_only_upload(template_bytes, [
        {"name": "TEST_NoPhone", "phone": None},
    ])
    files = {"file": ("leads.xlsx", payload,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = requests.post(f"{BASE_URL}/api/data-mgmt/import",
                      files=files, data={"mode": "upsert", "dry_run": "true"},
                      headers=headers, timeout=60)
    assert r.status_code == 200, r.text
    leads_ent = next(e for e in r.json()["entities"] if e["key"] == "leads")
    assert leads_ent["error"] == 1, leads_ent
    row_errs = leads_ent["rows"][0]["errors"]
    combined = " ".join(row_errs).lower()
    assert "hp" in combined or "phone" in combined, row_errs


# --- Full raw session ------------------------------------------------------
@pytest.fixture(scope="module")
def raw_session(headers):
    with open(RAW_FILE, "rb") as f:
        content = f.read()
    files = {"file": ("SIPRO_SemuaData.xlsx", content,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = requests.post(f"{BASE_URL}/api/data-mgmt/full/sessions",
                      files=files, headers=headers, timeout=90)
    assert r.status_code == 200, r.text
    return r.json()


def test_full_session_leads_parses_77(headers, raw_session):
    sid = raw_session["id"]
    r = requests.get(f"{BASE_URL}/api/data-mgmt/full/sessions/{sid}",
                     headers=headers, timeout=60)
    assert r.status_code == 200, r.text
    rep = r.json()
    leads_sheet = next((s for s in rep["sheets"] if s.get("collection") == "leads"), None)
    assert leads_sheet is not None, rep
    total = leads_sheet.get("total") or leads_sheet.get("rows") or 0
    assert total == 77, leads_sheet
    # No fatal error rows
    err = leads_sheet.get("error", 0)
    assert err == 0, f"parse errors: {leads_sheet}"


def test_full_session_commit_leads(headers, raw_session):
    sid = raw_session["id"]
    body = {"mode": "update", "sheets": ["leads"], "confirm": "IMPOR"}
    r = requests.post(f"{BASE_URL}/api/data-mgmt/full/sessions/{sid}/commit",
                      json=body, headers=headers, timeout=180)
    assert r.status_code == 200, r.text
    js = r.json()
    inserted = js["totals"].get("inserted", 0)
    assert inserted == 77, js["totals"]

    # Sample: fetch some leads by name and check defaults set (stage/assigned/score/+62)
    seen = []
    for nm in ("Nurmareta", "Fauzi"):
        r = requests.get(f"{BASE_URL}/api/leads?q={nm}&limit=10", headers=headers, timeout=30)
        assert r.status_code == 200
        body = r.json()
        items = body.get("data") or body.get("items") or []
        seen.extend([i for i in items if i.get("name") == nm])
    assert seen, "seeded leads not visible via /api/leads"
    for lead in seen:
        assert lead.get("stage") == "acquisition", lead
        assert (lead.get("phone") or "").startswith("+62"), lead
        assert lead.get("assigned_to"), lead
        assert lead.get("score") is not None, lead


# --- Cleanup ---------------------------------------------------------------
def test_zzz_cleanup_leads(headers):
    """Remove test leads via mongo (direct) to keep demo clean."""
    import asyncio
    from motor.motor_asyncio import AsyncIOMotorClient
    from dotenv import dotenv_values

    env = dotenv_values("/app/backend/.env")
    mongo_url = env.get("MONGO_URL") or os.environ.get("MONGO_URL")
    db_name = env.get("DB_NAME") or os.environ.get("DB_NAME")
    assert mongo_url and db_name

    async def _run():
        cli = AsyncIOMotorClient(mongo_url)
        d = cli[db_name]
        r1 = await d.leads.delete_many({"name": {"$in": TEST_NAMES}})
        # Raw import leads: created_by superadmin, source=import, email null
        r2 = await d.leads.delete_many({"created_by": SUPERADMIN[0],
                                        "source": "import",
                                        "email": None})
        cli.close()
        return r1.deleted_count, r2.deleted_count

    n1, n2 = asyncio.new_event_loop().run_until_complete(_run())
    print(f"cleanup: removed {n1} template leads, {n2} raw-import leads")
